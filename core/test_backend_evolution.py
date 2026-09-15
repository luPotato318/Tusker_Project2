import json
from datetime import timedelta
from io import BytesIO
from unittest.mock import patch
from urllib.parse import parse_qs

from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from .models import (
    Attendance,
    ChallengeSubmission,
    ChatMessage,
    ChatSession,
    JobApplication,
    JobOpportunity,
    PracticalChallenge,
    School,
    SchoolClass,
    SoftSkillAssessment,
    StudentProject,
    User,
    Workshop,
    WorkshopEnrollment,
)
from .risk import attendance_risk
from .services import employability_score, ranked_students
from .tutor import get_workshop_recommendations
from .views import _students_for_staff


class BackendEvolutionTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.school = School.objects.create(nome="Escola Horizonte", codigo="horizonte")
        cls.other_school = School.objects.create(nome="Escola Outra", codigo="outra")
        cls.classroom = SchoolClass.objects.create(escola=cls.school, nome="3A")
        cls.teacher = User.objects.create_user(
            "teacher@example.com", nome="Professora", perfil_acesso=User.Role.TEACHER, escola=cls.school,
        )
        cls.classroom.professores.add(cls.teacher)
        cls.student = User.objects.create_user(
            "student@example.com", nome="Ana", escola=cls.school, turma=cls.classroom,
        )
        cls.other_student = User.objects.create_user("other@example.com", nome="Bia", escola=cls.other_school)

    def test_ranking_queries_remain_constant_and_reuse_attendance_data(self):
        for student in (self.student, self.other_student):
            Attendance.objects.create(user=student, presente=True)
            Attendance.objects.create(user=student, data=timezone.localdate() - timedelta(days=1), presente=False)
            SoftSkillAssessment.objects.create(
                aluno=student, avaliador=self.teacher, comunicacao=3, proatividade=4, trabalho_equipe=5,
            )
            challenge = PracticalChallenge.objects.create(
                titulo="Entrega", descricao="Evidência", area="Educação", pontos=100, criado_por=self.teacher,
            )
            ChallengeSubmission.objects.create(
                aluno=student, desafio=challenge, resposta="Pronto", pontos_atribuidos=50,
                status=ChallengeSubmission.Status.APPROVED,
            )

        for count in (1, 2):
            with self.subTest(student_count=count), self.assertNumQueries(4):
                ranking = ranked_students(User.objects.filter(perfil_acesso=User.Role.STUDENT).order_by("pk")[:count])
                for item in ranking:
                    risk = attendance_risk(item["student"].frequencias)
                    self.assertEqual(risk["absence_rate"], 50.0)
                    self.assertIn(item["student"].escola.nome, {"Escola Horizonte", "Escola Outra"})
            self.assertEqual(len(ranking), count)
            self.assertEqual(ranking[0]["score"], 590)
            self.assertEqual(ranking[0]["components"], {"frequencia": 175, "soft_skills": 240, "entregas": 175})

        self.assertEqual(employability_score(self.student)["score"], ranking[0]["score"])

    def test_unassigned_teacher_cannot_access_unassigned_students(self):
        teacher = User.objects.create_user("unassigned-teacher@example.com", nome="Professor", perfil_acesso=User.Role.TEACHER)
        student = User.objects.create_user("unassigned-student@example.com", nome="Aluno sem escola")
        self.assertFalse(_students_for_staff(teacher).exists())
        self.client.force_login(teacher)
        response = self.client.post(reverse("soft_skill_assess", args=[student.pk]), {
            "comunicacao": 5, "proatividade": 5, "trabalho_equipe": 5,
        })
        self.assertEqual(response.status_code, 403)
        self.assertFalse(SoftSkillAssessment.objects.exists())

    def test_staff_scope_preserves_assigned_class_and_same_school_without_class(self):
        student_without_class = User.objects.create_user("no-class@example.com", nome="Sem turma", escola=self.school)
        self.assertSetEqual(set(_students_for_staff(self.teacher).values_list("pk", flat=True)), {
            self.student.pk, student_without_class.pk,
        })

    def test_invalid_tutor_payloads_do_not_write_or_call_provider(self):
        self.client.force_login(self.student)
        payloads = [[], None, True, 42, {}, {"mensagem": "   "}, {"mensagem": 5},
                    {"mensagem": "x" * 4001}, {"mensagem": "Olá", "session_id": True},
                    {"mensagem": "Olá", "session_id": "1"}, {"mensagem": "Olá", "session_id": 2 ** 64}]
        with patch("core.views.tutor_reply") as reply:
            for payload in payloads:
                with self.subTest(payload_type=type(payload).__name__):
                    response = self.client.post(reverse("tutor_api"), json.dumps(payload), content_type="application/json")
                    self.assertEqual(response.status_code, 400)
            reply.assert_not_called()
        self.assertFalse(ChatSession.objects.exists())
        self.assertFalse(ChatMessage.objects.exists())

    def test_tutor_continues_owned_session_and_rejects_foreign_session(self):
        session = ChatSession.objects.create(user=self.student)
        foreign_session = ChatSession.objects.create(user=self.other_student)
        self.client.force_login(self.student)
        with patch("core.views.tutor_reply", return_value={"texto": "Próximo passo", "provider": "local"}) as reply:
            denied = self.client.post(reverse("tutor_api"), {
                "mensagem": "Olá", "session_id": foreign_session.pk,
            }, content_type="application/json")
            self.assertEqual(denied.status_code, 404)
            reply.assert_not_called()
            accepted = self.client.post(reverse("tutor_api"), {
                "mensagem": "  Olá  ", "session_id": session.pk,
            }, content_type="application/json")
        self.assertEqual(accepted.status_code, 200)
        self.assertEqual(accepted.json()["session_id"], session.pk)
        self.assertEqual(session.mensagens.count(), 2)
        self.assertEqual(session.mensagens.first().content, "Olá")
        self.assertFalse(foreign_session.mensagens.exists())
        self.assertEqual(ChatSession.objects.count(), 2)

    @override_settings(PHP_BRIDGE_SECRET="bridge-test-secret")
    def test_bridge_rejects_invalid_payload_shapes(self):
        with patch("core.views.tutor_reply") as reply:
            for payload in ([], None, {"mensagem": ""}, {"mensagem": "Olá", "area": []}):
                with self.subTest(payload=payload):
                    response = self.client.post(
                        reverse("php_bridge_tutor_api"), json.dumps(payload), content_type="application/json",
                        HTTP_X_PIEM_BRIDGE_SECRET="bridge-test-secret",
                    )
                    self.assertEqual(response.status_code, 400)
            reply.assert_not_called()

    def test_expired_job_is_hidden_and_does_not_accept_applications(self):
        expired = JobOpportunity.objects.create(
            titulo="Encerrada", empresa="Empresa", descricao="Teste", publicada_por=self.teacher,
            encerra_em=timezone.now() - timedelta(seconds=1),
        )
        open_job = JobOpportunity.objects.create(
            titulo="Aberta", empresa="Empresa", descricao="Teste", publicada_por=self.teacher,
        )
        self.client.force_login(self.student)
        dashboard = self.client.get(reverse("dashboard"))
        self.assertEqual(list(dashboard.context["jobs"]), [open_job])
        self.assertEqual(self.client.post(reverse("job_apply", args=[expired.pk])).status_code, 302)
        self.assertFalse(JobApplication.objects.exists())
        self.client.post(reverse("job_apply", args=[open_job.pk]))
        self.assertTrue(JobApplication.objects.filter(aluno=self.student, vaga=open_job).exists())

    def test_tutor_recommends_only_open_visible_workshops(self):
        defaults = {
            "area": "Educação", "descricao": "Oficina", "data": timezone.now() + timedelta(days=2),
        }
        visible = Workshop.objects.create(titulo="Oficina da minha escola", escola=self.school, **defaults)
        Workshop.objects.create(titulo="Outra escola", escola=self.other_school, **defaults)
        Workshop.objects.create(titulo="Rascunho", status=Workshop.Status.DRAFT, **defaults)
        Workshop.objects.create(titulo="Prazo encerrado", inscricoes_ate=timezone.now() - timedelta(seconds=1), **defaults)
        Workshop.objects.create(titulo="Passada", **{**defaults, "data": timezone.now() - timedelta(days=1)})
        full = Workshop.objects.create(titulo="Lotada", vagas=1, **defaults)
        WorkshopEnrollment.objects.create(workshop=full, user=self.other_student)
        response = get_workshop_recommendations("Educação", user=self.student)
        self.assertIn(visible.titulo, response["texto"])
        for title in ("Outra escola", "Rascunho", "Prazo encerrado", "Passada", "Lotada"):
            self.assertNotIn(title, response["texto"])
        self.assertNotIn(visible.titulo, get_workshop_recommendations("Educação")["texto"])

    def test_registration_rejects_existing_contact_email_and_malformed_email(self):
        User.objects.create_user("12345678901", email="contact@example.com", nome="Existente")
        before = User.objects.count()
        for identifier in ("CONTACT@example.com", "invalid@"):
            with self.subTest(identifier=identifier):
                response = self.client.post(reverse("register"), {
                    "nome": "Novo", "identifier": identifier, "perfil_acesso": User.Role.STUDENT,
                    "senha": "uma-senha-segura",
                })
                self.assertEqual(response.status_code, 200)
                self.assertIn("identifier", response.context["form"].errors)
        self.assertEqual(User.objects.count(), before)

    def test_xlsx_export_writes_student_name_as_text(self):
        from openpyxl import load_workbook

        self.student.nome = "=1+1"
        self.student.save(update_fields=["nome"])
        admin = User.objects.create_superuser("admin@example.com", "test-password", nome="Gestora")
        self.client.force_login(admin)
        response = self.client.get(reverse("impact_export", args=["xlsx"]))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        names = [row[1] for row in workbook.active.iter_rows(min_row=2)]
        matching = next(cell for cell in names if cell.value == "=1+1")
        self.assertEqual(matching.data_type, "s")

    def test_resume_treats_user_markup_as_literal_text(self):
        self.student.nome = "Ana <b>"
        self.student.biografia = '<img src="arquivo-inexistente.png"/>'
        self.student.save(update_fields=["nome", "biografia"])
        StudentProject.objects.create(
            aluno=self.student, titulo="Projeto <b>", resumo="Texto com <i> aberto", area="Educação",
        )
        self.client.force_login(self.student)
        response = self.client.get(reverse("resume_pdf"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF-"))

    def test_course_pagination_preserves_filters_and_avoids_duplicate_results(self):
        event_date = timezone.now() + timedelta(days=3)
        Workshop.objects.bulk_create([
            Workshop(titulo=f"Oficina {number}", descricao="Prática", area="Educação", data=event_date)
            for number in range(13)
        ])
        Workshop.objects.create(titulo="Outro tema", descricao="Teste", area="Engenharia", data=event_date)
        filters = {"q": "Oficina", "area": "Educação", "modalidade": Workshop.Modality.ONSITE}
        first = self.client.get(reverse("courses"), filters)
        second = self.client.get(reverse("courses"), {**filters, "page": "2"})
        self.assertEqual(first.context["paginator"].count, 13)
        self.assertEqual(len(first.context["courses"]), 12)
        self.assertEqual(len(second.context["courses"]), 1)
        self.assertTrue(set(first.context["courses"]).isdisjoint(second.context["courses"]))
        self.assertEqual(parse_qs(second.context["pagination_query"]), {key: [value] for key, value in filters.items()})
        self.assertContains(second, "Página 2 de 2")
        invalid_page = self.client.get(reverse("courses"), {**filters, "page": "inválida"})
        self.assertEqual(invalid_page.context["page_obj"].number, 1)
